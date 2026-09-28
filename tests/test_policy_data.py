"""Tests for policy_data.py: Eurostat / OECD export parsers, the GBARD
splice, reform-event detection and consolidation episodes."""
import numpy as np
import openpyxl
import pandas as pd
import pytest

import policy_data as P


def _eurostat_xlsx(path, sheets):
    """Write a workbook laid out like a Eurostat 'spreadsheet' export.
    sheets: list of (meta dict, years, {country: [(value, flag), ...]})."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for k, (meta, years, data) in enumerate(sheets, 1):
        ws = wb.create_sheet(f"Sheet {k}")
        ws.append(["Data extracted on ", None, None])
        ws.append(["Dataset: ", "test", None])
        ws.append([None])
        for key, val in meta.items():
            ws.append([key, None, val])
        ws.append([None])
        row = ["TIME"]
        for y in years:
            row += [str(y), None]
        ws.append(row)
        ws.append(["GEO (Labels)"])
        for c, vals in data.items():
            r = [c]
            for v, f in vals:
                r += [v, f]
            ws.append(r)
        ws.append(["Special value"])
        ws.append([":", "not available"])
    wb.save(path)


def test_read_eurostat_picks_sheet_by_metadata(tmp_path):
    f = tmp_path / "g.xlsx"
    _eurostat_xlsx(f, [
        ({"Nomenclature": "Defence", "Unit of measure": "Million euro"},
         [2005, 2006], {"Germany (until 1990 former territory of the FRG)":
                        [(1.0, None), (2.0, None)]}),
        ({"Nomenclature": "Total appropriations",
          "Unit of measure": "Million euro"},
         [2005, 2006], {"Germany (until 1990 former territory of the FRG)":
                        [(10.0, "e"), (":", None)],
                        "Slovak Republic": [(3.0, None), (4.0, None)]})])
    d = P.read_eurostat(str(f), **{"unit of measure": "Million euro",
                                   "nomenclature": "Total appropriations"})
    de = d[d.Country == "Germany"].set_index("Year")
    assert de.loc[2005, "value"] == 10 and de.loc[2005, "flag"] == "e"
    assert np.isnan(de.loc[2006, "value"])                  # ':' → missing
    assert "Slovakia" in set(d.Country)                     # name normalised
    with pytest.raises(ValueError):
        P.read_eurostat(str(f), nomenclature="Nope")


def test_gbard_splice_prefers_nabs2007(tmp_path):
    pdir = tmp_path / "policy"
    pdir.mkdir()
    meta92 = {"Nomenclature": "Total appropriations",
              "Unit of measure": "Million euro"}
    meta07 = {"Nomenclature": "Total government budget allocations for R&D",
              "Unit of measure": "Million euro"}
    _eurostat_xlsx(pdir / P.FILES["gba92"], [(meta92, [2006, 2007],
                   {"Austria": [(100.0, None), (110.0, None)]})])
    _eurostat_xlsx(pdir / P.FILES["gba07"], [(meta07, [2007, 2008],
                   {"Austria": [(111.0, None), (120.0, None)]})])
    d = P.build(str(pdir)).set_index("Year")
    assert d.loc[2006, "GBARD_mEUR"] == 100 and d.loc[2006, "GBARD_src"] == "NABS1992"
    assert d.loc[2007, "GBARD_mEUR"] == 111 and d.loc[2007, "GBARD_src"] == "NABS2007"
    assert d.loc[2008, "GBARD_mEUR"] == 120


def test_read_oecd_rdsub(tmp_path):
    f = tmp_path / "o.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, "Implied tax subsidy rates"])
    ws.append([None])
    ws.append([None, "Time period", "Time period", "Time period", "2000", "2001"])
    ws.append([None, "Reference area", "Firm size"])
    ws.append([None, "Profit scenario: Profitable"])
    ws.append([None, "·  Malta", "SME", None, 0.1, 0.2])
    ws.append([None, "Slovak Republic", "Large firm", None, 0.0, 0.3])
    ws.append([None, "Profit scenario: Loss-making"])
    ws.append([None, "Slovak Republic", "Large firm", None, -0.01, 0.1])
    wb.save(f)
    d = P.read_oecd_rdsub(str(f))
    sk = d[(d.Country == "Slovakia") & (d.scenario == "profit")].set_index("Year")
    assert sk.loc[2001, "value"] == 0.3
    assert set(d.Country) == {"Malta", "Slovakia"}
    assert set(d.scenario) == {"profit", "loss"}


def test_tax_reform_events_require_persistence():
    rows = []
    for y in range(2000, 2012):
        rows.append(("A", y, 0.30 if y >= 2005 else 0.05))      # lasting
        rows.append(("B", y, 0.30 if y == 2006 else 0.05))      # one-off
        rows.append(("C", y, 0.05))                              # none
    df = pd.DataFrame(rows, columns=["Country", "Year",
                                     "RD_subsidy_large_profit"])
    assert P.tax_reform_events(df) == {"A": 2005}


def test_consolidation_episodes():
    df = pd.DataFrame({"Country": ["A"] * 4, "Year": [1, 2, 3, 4],
                       "Gov_balance_pct_GDP": [-5.0, -3.0, -2.5, -0.5]})
    np.testing.assert_array_equal(P.consolidation_episodes(df).values,
                                  [np.nan, 1, 0, 1])


def test_coverage_reports_gaps():
    df = pd.DataFrame({"Country": ["A"] * 3, "Year": [2000, 2002, 2003],
                       "x": [1.0, 2.0, 3.0]})
    cov = P.coverage(df, ["A", "B"], ["x"], window=(2000, 2003))
    assert cov.set_index("Country").loc["A", "x"] == "2000–2003 (1 gaps)"
    assert cov.set_index("Country").loc["B", "x"] == "—"
